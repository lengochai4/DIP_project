"""Build the editable Vietnamese report from frozen G7 records and Markdown.

Run with the primary-runtime Python. No project algorithm or evidence is changed.
"""

from pathlib import Path
import csv
import hashlib
import json
import re
import subprocess
import sys
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.graphics.shapes import Drawing, Rect, String, Line, Polygon
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics import renderPDF
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
OUTPUT = ROOT / "output/documents/Bao_cao_tam_DIP_Touchless_STEM.docx"
TMP = ROOT / "runs/report-temp"
ASSETS = HERE / "assets"
RUNTIME = Path(r"C:\Users\DIABLO\.cache\codex-runtimes\codex-primary-runtime")
POPPLER = RUNTIME / "dependencies/native/poppler/Library/bin/pdftoppm.exe"
TITLE_TEXT = "Xây dựng hệ thống tương tác không chạm thời gian thực cho học tập STEM 3D sử dụng tiền xử lý ảnh thích ứng và lọc tín hiệu chuyển động từ webcam RGB"
TITLE = TITLE_TEXT.upper()
FONT = "Times New Roman"
BODY_PATH = HERE / "BAO_CAO_TAM.md"
COVER_STATUS = "BẢN TẠM"
RESULT_SUMMARY = "Hai lượt A1 có dữ liệu cho thấy F1 và F2 giảm độ rung so với F0, với kết quả rất gần nhau. A2 chưa tính được chỉ số chính. P0/P1 có tỷ lệ hợp lệ bằng nhau và CLAHE không kích hoạt trong các cửa sổ B. Bản minh họa G7 đã chạy tới cảnh 3D, còn hạn chế nhận dạng ngoài ý muốn."


def records(group):
    with (ROOT / f"submission/evidence/{group}/metrics.csv").open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def validate_evidence():
    a1, a2, bn, bl = map(records, ("a1_static", "a2_dynamic", "b_normal", "b_lowlight"))
    assert len(a1) == len(a2) == 9 and len(bn) == len(bl) == 6
    assert {(r["trial_id"], r["profile"]) for r in a1 if r["available"] == "false"} == {
        ("trial-002", f) for f in ("F0", "F1", "F2")
    }
    assert all(r["available"] == "false" and r["value"] == "" for r in a2)
    assert all(r["unavailable_reason"] == "no_common_usable_frames" for r in a2)
    assert all(int(r["sample_count"]) == 241 for r in bn + bl)
    assert all(float(r["value"]) == 0 for r in bl)
    for trial in ("trial-001", "trial-002", "trial-003"):
        values = [float(r["value"]) for r in bn if r["trial_id"] == trial]
        assert values[0] == values[1]
    text = (HERE / "BAO_CAO_TAM.md").read_text(encoding="utf-8")
    for r in a1:
        if r["available"] == "true":
            assert f"{float(r['value']):.9f}".replace(".", ",") in text
    for r in bn:
        if float(r["value"]):
            assert f"{float(r['value']) * 100:.6f}%".replace(".", ",") in text
    # The supplied sample is a layout reference, not the project's identity/results.
    for unwanted in ("SV2026-468", "Trần Gia Kiệt", "Trần Quang Khải", "MLP–BiLSTM", "97.20%", "24110103"):
        assert unwanted not in text
    return a1, bn


def save_drawing(drawing, name):
    path = TMP / f"{name}.pdf"
    renderPDF.drawToFile(drawing, str(path))
    subprocess.run([str(POPPLER), "-singlefile", "-r", "180", "-png", str(path), str(ASSETS / name)], check=True)


def make_assets(a1, bn):
    ASSETS.mkdir(parents=True, exist_ok=True)
    TMP.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(TTFont("ReportTimes", r"C:\Windows\Fonts\times.ttf"))
    d = Drawing(470, 260)
    labels = [
        ("Webcam hoặc", "nguồn phát lại"), ("Quản lý ROI", "giữ hình học ảnh"),
        ("Ánh sáng HSV", "CLAHE hoặc bỏ qua"), ("MediaPipe", "điểm mốc bàn tay"),
        ("InteractionState", "lệnh độc lập cảnh"), ("Ánh xạ cử chỉ", "xoay và kích thước"),
        ("Lọc thời gian", "F0 / F1 / F2"), ("Kiểm tra", "quan sát hợp lệ"),
    ]
    centers = [(59 + 117 * i, 222) for i in range(4)] + [(59 + 117 * i, 135) for i in range(4)]
    for (x, y), lines in zip(centers, labels):
        d.add(Rect(x - 53, y - 25, 106, 50, fillColor=None, strokeWidth=.7))
        for j, label in enumerate(lines):
            d.add(String(x, y + 6 - 17 * j, label, fontName="ReportTimes", fontSize=12, textAnchor="middle"))

    def arrow(x1, y1, x2, y2):
        d.add(Line(x1, y1, x2, y2, strokeWidth=.9))
        if x1 == x2:
            d.add(Polygon([x2, y2, x2-3, y2+6, x2+3, y2+6], fillColor=None))
        else:
            sign = 1 if x2 > x1 else -1
            d.add(Polygon([x2, y2, x2-sign*6, y2-3, x2-sign*6, y2+3], fillColor=None))
    for i in range(3):
        arrow(centers[i][0]+53, 222, centers[i+1][0]-53, 222)
        arrow(centers[7-i][0]-53, 135, centers[6-i][0]+53, 135)
    arrow(410, 197, 410, 160)
    arrow(59, 110, 59, 70)
    d.add(Rect(6, 20, 106, 50, fillColor=None, strokeWidth=.7))
    d.add(String(59, 50, "Cảnh 3D", fontName="ReportTimes", fontSize=12, textAnchor="middle"))
    d.add(String(59, 33, "minh họa", fontName="ReportTimes", fontSize=12, textAnchor="middle"))
    d.add(String(260, 60, "TrackingFrame và nhật ký", fontName="ReportTimes", fontSize=12, textAnchor="middle"))
    d.add(String(260, 40, "được lưu để phân tích thực nghiệm", fontName="ReportTimes", fontSize=12, textAnchor="middle"))
    save_drawing(d, "pipeline")

    from reportlab.lib.colors import HexColor
    for name, rows, profiles, trials, scale, maximum, title in (
        ("a1_jitter", a1, ("F0", "F1", "F2"), ("trial-001", "trial-003"), 1, .035, "Radial RMS jitter trong tọa độ chuẩn hóa"),
        ("b_normal", bn, ("P0", "P1"), ("trial-001", "trial-002", "trial-003"), 100, 60, "Tỷ lệ khung có quan sát bàn tay hợp lệ (%)"),
    ):
        d = Drawing(470, 280)
        chart = VerticalBarChart()
        chart.x, chart.y, chart.width, chart.height = 58, 47, 392, 184
        chart.data = [[next(float(r["value"]) * scale for r in rows if r["trial_id"] == trial and r["profile"] == p) for trial in trials] for p in profiles]
        chart.categoryAxis.categoryNames = list(trials)
        chart.categoryAxis.labels.fontName = chart.valueAxis.labels.fontName = "ReportTimes"
        chart.categoryAxis.labels.fontSize = chart.valueAxis.labels.fontSize = 12
        chart.valueAxis.valueMin, chart.valueAxis.valueMax = 0, maximum
        chart.valueAxis.valueStep = .005 if scale == 1 else 10
        chart.barSpacing, chart.groupSpacing = 4, 25
        for i, p in enumerate(profiles):
            chart.bars[i].fillColor = HexColor(("#333333", "#888888", "#CCCCCC")[i])
            chart.bars[i].strokeColor = HexColor("#222222")
            d.add(Rect(125 + 80*i, 10, 12, 10, fillColor=chart.bars[i].fillColor, strokeWidth=.5))
            d.add(String(143 + 80*i, 10, p, fontName="ReportTimes", fontSize=12))
        d.add(chart)
        d.add(String(235, 260, title, fontName="ReportTimes", fontSize=14, textAnchor="middle"))
        if name == "a1_jitter":
            d.add(String(235, 241, "trial-001: n = 5; trial-003: n = 128; trial-002 chưa tính được", fontName="ReportTimes", fontSize=10, textAnchor="middle"))
        else:
            d.add(String(235, 241, "241 khung / lượt / điều kiện; P1 không kích hoạt CLAHE", fontName="ReportTimes", fontSize=10, textAnchor="middle"))
        save_drawing(d, name)


def fmt_run(run, size=14, bold=False):
    run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().rFonts
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + key), None)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), FONT)


def field(paragraph, instruction, cached="1"):
    for kind, text in (("begin", None), (None, instruction), ("separate", None), (None, cached), ("end", None)):
        r = OxmlElement("w:r")
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        elif text == instruction:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        else:
            element = OxmlElement("w:t")
            element.text = text
        r.append(element)
        paragraph._p.append(r)


def paragraph(doc, text, *, bold=False, align=None, size=14, after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if align is None else align
    p.paragraph_format.space_after = Pt(after)
    fmt_run(p.add_run(text), size, bold)
    return p


def table(doc, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = widths or [15.8 / len(rows[0])] * len(rows[0])
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for i, row in enumerate(rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for j, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 or len(str(value)) > 20 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            if len(rows) <= 5:
                p.paragraph_format.keep_with_next = i < len(rows) - 1
            fmt_run(p.add_run(str(value)), 11.5, i == 0)
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement("w:" + edge)
                for k, v in (("val", "single"), ("sz", "5"), ("color", "D9D9D9")):
                    e.set(qn("w:" + k), v)
                borders.append(e)
            pr.append(borders)
            margin = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement("w:" + edge)
                e.set(qn("w:w"), "75")
                e.set(qn("w:type"), "dxa")
                margin.append(e)
            pr.append(margin)
            if i == 0:
                sh = OxmlElement("w:shd")
                sh.set(qn("w:fill"), "E8E8E8")
                pr.append(sh)
        trpr = t.rows[i]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def configure_section(section):
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.0)
    section.left_margin, section.right_margin = Cm(3.0), Cm(2.0)
    section.footer_distance = Cm(1.0)


def cover(doc, inner=False):
    center = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc, "BỘ GIÁO DỤC VÀ ĐÀO TẠO", align=center, size=14, after=8)
    paragraph(doc, "TRƯỜNG ................................................................", align=center, bold=True, size=14, after=65)
    paragraph(doc, "BÁO CÁO TỔNG KẾT", align=center, size=19, bold=True, after=26)
    paragraph(doc, "ĐỀ TÀI NGHIÊN CỨU KHOA HỌC CỦA SINH VIÊN", align=center, size=15, bold=True, after=38)
    p = paragraph(doc, TITLE, align=center, size=16, bold=True, after=20)
    p.style = doc.styles["Title"]
    p.paragraph_format.line_spacing = 1.25
    paragraph(doc, "Mã đề tài: ........................................................", align=center, after=20)
    if inner:
        for line in (
            "Thuộc nhóm ngành khoa học: ........................................",
            "SV thực hiện: .......................................  Nam, Nữ: ..........",
            "Dân tộc: ........................................................................",
            "Lớp, khoa: ....................................................................",
            "Năm thứ: ...................  Số năm đào tạo: ...........................",
            "Ngành học: ..................................................................",
            "Người hướng dẫn: ........................................................",
        ):
            p = paragraph(doc, line, align=WD_ALIGN_PARAGRAPH.LEFT, size=13, after=5)
            p.paragraph_format.line_spacing = 1.1
    else:
        paragraph(doc, "Chủ nhiệm đề tài: ...................................................", align=center, after=32)
        paragraph(doc, COVER_STATUS, align=center, size=13, after=45)
    paragraph(doc, "........................, tháng 10 năm 2026", align=center, bold=True, after=0)


def build_docx(pages=None):
    pages = pages or {}
    doc = Document()
    doc.core_properties.title = "Báo cáo tạm về hệ thống tương tác STEM từ webcam RGB"
    doc.core_properties.author = ""
    doc.core_properties.subject = "Xử lý ảnh và lọc tín hiệu chuyển động"
    for section in doc.sections:
        configure_section(section)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        st = doc.styles[name]
        for border in list(st.element.findall(".//" + qn("w:pBdr"))):
            border.getparent().remove(border)
        st.font.name, st.font.size = FONT, Pt(14)
        st.font.color.rgb = RGBColor(0, 0, 0)
        fonts = st.element.get_or_add_rPr().rFonts
        for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + key), None)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + key), FONT)
        st.paragraph_format.line_spacing = 1.25
        st.paragraph_format.space_after = Pt(4)
    for level, size in ((1, 16), (2, 14), (3, 14)):
        st = doc.styles[f"Heading {level}"]
        st.font.size, st.font.bold = Pt(size), True
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(12 if level > 1 else 0)
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for edge, space in (("top", "30"), ("bottom", "30"), ("left", "55"), ("right", "35")):
        e = OxmlElement("w:" + edge)
        for key, value in (("val", "double"), ("sz", "6"), ("space", space), ("color", "000000")):
            e.set(qn("w:" + key), value)
        borders.append(e)
    doc.sections[0]._sectPr.append(borders)
    cover(doc)
    doc.add_page_break()
    cover(doc, True)
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(section)
    pr = section._sectPr.find(qn("w:pgBorders"))
    if pr is not None:
        section._sectPr.remove(pr)
    paragraph(doc, "MỤC LỤC", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, after=16)
    raw = BODY_PATH.read_text(encoding="utf-8")
    headings = [(len(m.group(1)), m.group(2)) for m in re.finditer(r"^(#{1,3}) (.+)$", raw, re.M)]
    extras = [(1, "DANH MỤC BẢNG BIỂU"), (1, "DANH MỤC TỪ VIẾT TẮT"), (1, "THÔNG TIN KẾT QUẢ NGHIÊN CỨU CỦA ĐỀ TÀI")]
    toc_entries = extras + headings
    toc_paragraphs = []
    for level, text in toc_entries:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.08
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm((level - 1) * .45)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.8 - (level-1)*.45), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        fmt_run(p.add_run(text + "\t" + str(pages.get(text, "…"))), 12, level == 1)
        toc_paragraphs.append(p)
    # A native TOC field surrounds its current display paragraphs; Word can update it.
    start = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); start.append(begin)
    code = OxmlElement("w:r"); instr = OxmlElement("w:instrText"); instr.text = ' TOC \\o "1-3" \\h \\z \\u '; code.append(instr)
    separate = OxmlElement("w:r"); sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate"); separate.append(sep)
    for offset, element in enumerate((start, code, separate)):
        toc_paragraphs[0]._p.insert(offset, element)
    end = OxmlElement("w:r"); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); end.append(e); toc_paragraphs[-1]._p.append(end)
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(section)
    section.footer.is_linked_to_previous = False
    num = OxmlElement("w:pgNumType"); num.set(qn("w:start"), "1"); section._sectPr.append(num)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(footer, " PAGE ")
    for r in footer.runs:
        fmt_run(r, 12)
    doc.add_heading(extras[0][1], 1)
    captions = re.findall(r"^Bảng .+$", raw, re.M) + re.findall(r"!\[([^]]+)\]", raw)
    for caption in captions:
        p = paragraph(doc, caption, size=12, align=WD_ALIGN_PARAGRAPH.LEFT, after=8)
    doc.add_page_break()
    doc.add_heading(extras[1][1], 1)
    table(doc, [
        ["Từ viết tắt", "Diễn giải", "Ý nghĩa"],
        ["DIP", "Digital Image Processing", "Xử lý ảnh số"],
        ["STEM", "Science Technology Engineering Mathematics", "Khoa học, công nghệ, kỹ thuật và toán học"],
        ["RGB / BGR", "Thứ tự các kênh màu", "Đỏ, lục, lam / lam, lục, đỏ"],
        ["HSV", "Hue Saturation Value", "Sắc độ, bão hòa, độ sáng"],
        ["ROI", "Region of Interest", "Vùng quan tâm"],
        ["CLAHE", "Contrast Limited Adaptive Histogram Equalization", "Cân bằng lược đồ thích ứng có giới hạn tương phản"],
        ["EMA", "Exponential Moving Average", "Trung bình trượt lũy thừa"],
        ["RMS / RMSE", "Root Mean Square / Root Mean Square Error", "Căn trung bình bình phương / căn sai số bình phương trung bình"],
        ["F0 / F1 / F2", "Các đường lọc", "Thô / 1-Euro cố định / thích ứng có giới hạn"],
        ["P0 / P1", "Các đường tiền xử lý", "Bỏ qua CLAHE / CLAHE thích ứng"],
        ["FPS", "Frames Per Second", "Số khung hình mỗi giây"],
        ["SHA-256", "Hàm băm", "Đối chiếu định danh tài nguyên"],
    ], [2.5, 5.2, 8.1])
    doc.add_page_break()
    doc.add_heading(extras[2][1], 1)
    paragraph(doc, "1. Thông tin chung", bold=True)
    for s in (
        "Tên đề tài: " + TITLE_TEXT,
        "Chủ nhiệm đề tài: ................................  Mã số SV: .....................",
        "Lớp: ...................................  Khoa: ............................................",
        "Người hướng dẫn: ..........................................................................",
    ):
        paragraph(doc, s, align=WD_ALIGN_PARAGRAPH.LEFT)
    paragraph(doc, "Thành viên đề tài", bold=True)
    table(doc, [["STT", "Họ và tên", "MSSV", "Lớp", "Khoa"], ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""]], [1.3, 4.5, 2.5, 3, 3.5])
    for title, body in (
        ("2. Mục tiêu đề tài", "Xây dựng nền xử lý ảnh và lọc chuyển động từ một webcam RGB, hỗ trợ tương tác với cảnh STEM 3D và đánh giá bằng nguồn phát lại có định danh."),
        ("3. Nội dung thực hiện và đóng góp", "Triển khai quản lý ROI, quyết định ánh sáng, chính sách CLAHE, các đường lọc F0/F1/F2 và trạng thái tương tác độc lập với cảnh. Bộ lọc F2 có giới hạn beta và tần số cắt; các phương pháp nền MediaPipe, CLAHE và 1-Euro được kế thừa từ tài liệu."),
        ("4. Kết quả nghiên cứu", RESULT_SUMMARY),
        ("5. Khả năng áp dụng", "Nền xử lý có thể phục vụ thực hành về xử lý ảnh và tín hiệu chuyển động. Việc sử dụng trong học tập cần được đánh giá thêm theo nhiệm vụ, điều kiện camera và thiết kế giao diện cụ thể; báo cáo chưa chứng minh hiệu quả học tập."),
        ("6. Công bố khoa học hoặc nhận xét của cơ sở áp dụng", "..........................................................................................................."),
    ):
        paragraph(doc, title, bold=True).paragraph_format.keep_with_next = True
        paragraph(doc, body)
    paragraph(doc, "Ngày ...... tháng ...... năm ......", align=WD_ALIGN_PARAGRAPH.RIGHT)
    paragraph(doc, "SV chịu trách nhiệm chính thực hiện đề tài\n(Ký, họ và tên)", align=WD_ALIGN_PARAGRAPH.RIGHT, after=35)
    paragraph(doc, "Nhận xét của người hướng dẫn", bold=True)
    for _ in range(3):
        paragraph(doc, "." * 105, after=2)
    paragraph(doc, "Ngày ...... tháng ...... năm ......\nNgười hướng dẫn\n(Ký, họ và tên)", align=WD_ALIGN_PARAGRAPH.RIGHT)
    lines = raw.splitlines()
    i = 0
    previous_major = None
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line.startswith("#"):
            level = len(line.split(" ", 1)[0]); text = line.split(" ", 1)[1]
            if level == 1 and previous_major != "PHẦN NỘI DUNG":
                doc.add_page_break()
            p = doc.add_heading(text, level)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
            if level == 1:
                previous_major = text
            i += 1; continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", v) for v in row):
                    rows.append(row)
                i += 1
            n = len(rows[0])
            widths = [2.4, 3.6, 3.3, 3.3, 3.2] if n == 5 else [3.1, 4.2, 4.2, 4.3] if n == 4 else [3.0, 7.0, 5.8] if n == 3 else [5.0, 10.8]
            table(doc, rows, widths)
            continue
        match = re.fullmatch(r"!\[(.+)\]\((.+)\)", line)
        if match:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(HERE / match.group(2)), width=Cm(15.4))
            p = paragraph(doc, match.group(1), align=WD_ALIGN_PARAGRAPH.CENTER, size=12, after=10)
            i += 1; continue
        if line.startswith("Bảng "):
            p = paragraph(doc, line, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.keep_with_next = True
            i += 1; continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if line.startswith("> Tỷ lệ"):
                fmt_run(p.add_run(line[2:]), 12)
                i += 1; continue
            math = OxmlElement("m:oMath")
            identifiers = {"fc,max": ("f", "c,max"), "fc,min": ("f", "c,min"), "fbase": ("f", "base"), "fmin": ("f", "min"), "fc": ("f", "c"), "vmax": ("v", "max"), "vβ": ("v", "β"), "gv": ("g", "v"), "αd": ("α", "d"), "βbase": ("β", "base"), "βmin": ("β", "min"), "βmax": ("β", "max")}
            def math_run(value):
                r = OxmlElement("m:r"); t = OxmlElement("m:t"); t.text = value; r.append(t)
                return r
            pattern = "(" + "|".join(re.escape(k) for k in identifiers) + ")"
            for token in re.split(pattern, line[2:]):
                if token in identifiers:
                    sub = OxmlElement("m:sSub")
                    base = OxmlElement("m:e"); base.append(math_run(identifiers[token][0]))
                    lower = OxmlElement("m:sub"); lower.append(math_run(identifiers[token][1]))
                    sub.append(base); sub.append(lower); math.append(sub)
                elif token:
                    math.append(math_run(token))
            p._p.append(math)
            p.paragraph_format.space_after = Pt(8)
            i += 1; continue
        if line.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                p = paragraph(doc, lines[i], align=WD_ALIGN_PARAGRAPH.LEFT, size=11, after=2)
                p.paragraph_format.keep_with_next = not lines[i+1].startswith("```")
                for r in p.runs:
                    r.font.name = "Consolas"
                i += 1
            i += 1; continue
        text = line
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "![", "> ", "```", "Bảng ")):
            text += " " + lines[i].strip(); i += 1
        p = paragraph(doc, text, size=12 if text.startswith("[") else 14,
                      align=WD_ALIGN_PARAGRAPH.LEFT if text.startswith("[") else None)
        if not text.startswith("["):
            p.paragraph_format.first_line_indent = Cm(.75)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    (TMP / "headings.json").write_text(json.dumps([t for _, t in toc_entries], ensure_ascii=False, indent=2), encoding="utf-8")
    return OUTPUT


def main():
    if len(sys.argv) > 1 and sys.argv[1] == "--pages":
        path = Path(sys.argv[2])
        reader = PdfReader(path)
        texts = [(p.extract_text() or "").splitlines() for p in reader.pages]
        headings = json.loads((TMP / "headings.json").read_text(encoding="utf-8"))
        def matches(lines, heading):
            normalized = " ".join(heading.split())
            return any(" ".join(" ".join(lines[j:j+n]).split()) == normalized
                       for j in range(len(lines)) for n in (1, 2, 3))
        start = next(i for i, lines in enumerate(texts) if matches(lines, "DANH MỤC BẢNG BIỂU"))
        mapping = {}
        for h in headings:
            found = [i for i, lines in enumerate(texts) if i >= start and matches(lines, h)]
            if found:
                mapping[h] = found[0] - start + 1
        (TMP / "page_map.json").write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding="utf-8")
        assert len(mapping) == len(headings), sorted(set(headings) - set(mapping))
        build_docx(mapping)
    else:
        a1, bn = validate_evidence()
        make_assets(a1, bn)
        build_docx()
        inputs = [p for directory in (ROOT / "submission/evidence", ROOT / "experiments/final") for p in directory.rglob("*") if p.is_file()]
        provenance = {
            "base_commit": "f454c6b8325c85199c0122c0e822fe8e76c1526c",
            "report_branch": "docs/report-temp-from-main",
            "reference_role": "layout only; all reference personal data and experiment results excluded",
            "g7_evidence_modified": False,
            "metrics": "frozen CSV values; B converted to percentages; unavailable values remain unavailable",
            "source_hashes": {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs},
        }
        (HERE / "provenance.json").write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUTPUT)


if __name__ == "__main__":
    main()
